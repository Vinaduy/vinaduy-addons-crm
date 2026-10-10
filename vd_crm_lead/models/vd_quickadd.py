# -*- coding: utf-8 -*-
"""THÊM KHÁCH HÀNG MỚI — popup OWL tự viết (user 2026-10-06).

Backend sạch, độc lập: nhận danh sách KH (gõ tay hoặc nạp từ Excel) dưới dạng
JSON từ frontend → chống trùng SĐT → parse cột "Thông tin" vào trường khai thác
→ tạo lead + chia cho nhân viên. KHÔNG dùng lại wizard transient cũ.
"""
import base64
import io
import logging
import re

from odoo import api, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# Nguồn khách chọn được trong popup.
VD_QA_SOURCES = [
    ('quet', 'Quét số'),
    ('zalo', 'Zalo'),
    ('tiktok', 'TikTok'),
    ('facebook', 'Facebook'),
    ('referral', 'Giới thiệu'),
    ('manual', 'Khác'),
]
# Nguồn → kênh báo cáo (vd_lead_channel).
VD_QA_CHANNEL = {'quet': 'quet', 'zalo': 'zalo', 'tiktok': 'tiktok', 'facebook': 'facebook'}

# Header nhận dạng cột khi nạp Excel.
def _novn(s):
    """Bỏ dấu tiếng Việt + chuẩn hoá dấu nhân, về lowercase (để dò từ khoá)."""
    s = (s or '')
    for a, b in (('áàảãạâấầẩẫậăắằẳẵặ', 'a'), ('éèẻẽẹêếềểễệ', 'e'),
                 ('íìỉĩị', 'i'), ('óòỏõọôốồổỗộơớờởỡợ', 'o'),
                 ('úùủũụưứừửữự', 'u'), ('ýỳỷỹỵ', 'y'), ('đ', 'd')):
        for ch in a:
            s = s.replace(ch, b).replace(ch.upper(), b)
    return s.replace('✕', 'x').replace('×', 'x').lower()


class VdQuickAddLead(models.Model):
    _inherit = 'crm.lead'

    # ============================================================= META
    @api.model
    def vd_quickadd_meta(self):
        """Nguồn + danh sách NV cho popup."""
        return {
            'sources': [{'key': k, 'label': l} for k, l in VD_QA_SOURCES],
            'nvs': self._vd_qa_nv_list(),
        }

    @api.model
    def _vd_qa_nv_list(self):
        """NV bán hàng thật (loại admin/share) — cho ô chọn chia số."""
        Users = self.env['res.users'].sudo()
        gid = self.env.ref('sales_team.group_sale_salesman').id
        nvs = Users.search([
            ('share', '=', False), ('active', '=', True), ('groups_id', 'in', gid),
        ]).filtered(lambda u: not (u._is_admin() or u.has_group('base.group_system')))
        return [{'id': u.id, 'name': u.name,
                 'team': (getattr(u, 'vd_team_label', '') or ''),
                 'on': bool(getattr(u, 'vd_can_receive_pancake', True))}
                for u in nvs.sorted('name')]

    @api.model
    def vd_quickadd_check_phone(self, phone):
        """Kiểm tra NGAY 1 SĐT (gõ tay) có TRÙNG trong hệ thống không — để báo
        trùng ngay khi thêm, nhất quán với lúc nạp Excel."""
        core = self._vd_qa_core(phone)
        if not core or not self._vd_qa_phone_ok('0' + core):
            return {'bad': True, 'exists': False}
        return {'bad': False, 'exists': bool(self._vd_qa_existing_cores([core]))}

    # ============================================================ HELPERS
    def _vd_qa_core(self, phone):
        s = self._vd_normalize_phones_set(phone)
        return next(iter(s)) if s else ''

    @staticmethod
    def _vd_qa_phone_ok(cphone):
        return bool(re.match(r'^0[35789]\d{8}$', cphone or ''))

    def _vd_qa_existing_cores(self, cores):
        """Set core-phone ĐÃ CÓ trong hệ thống (kể cả archive)."""
        if not cores:
            return set()
        variants = []
        for c in cores:
            variants += ['0' + c, c, '84' + c, '+84' + c]
        out = set()
        for r in self.with_context(active_test=False).sudo().search_read(
                [('phone', 'in', variants)], ['phone']):
            cc = self._vd_qa_core(r.get('phone'))
            if cc:
                out.add(cc)
        return out

    def _vd_qa_dedup(self, rows):
        """rows=[{phone,name,source,info,excel}] → (clean, dup, bad).
        Loại SĐT sai, trùng trong lô, và trùng số đã có trong hệ thống."""
        seen = set()
        staged = []
        dups = []
        bad = 0
        for r in rows:
            core = self._vd_qa_core(r.get('phone'))
            cphone = '0' + core if core else ''
            if not core or not self._vd_qa_phone_ok(cphone):
                bad += 1
                continue
            info = (r.get('info') or '').strip()
            if core in seen:
                dups.append({'core': core, 'info': info})
                continue
            seen.add(core)
            nm = (r.get('name') or '').strip() or ('Khách ' + cphone[-4:])
            staged.append({'phone': cphone, 'core': core, 'name': nm,
                           'source': r.get('source') or 'quet',
                           'info': info, 'excel': bool(r.get('excel'))})
        existing = self._vd_qa_existing_cores([s['core'] for s in staged])
        clean = []
        for s in staged:
            if s['core'] in existing:
                dups.append(s)
            else:
                clean.append(s)
        return clean, dups, bad

    def _vd_qa_find_lead_by_core(self, core):
        """Lead đang có cho 1 core SĐT — ưu tiên ĐANG CHĂM (active), mới nhất."""
        variants = ['0' + core, core, '84' + core, '+84' + core]
        leads = self.with_context(active_test=False).sudo().search(
            [('phone', 'in', variants)], order='active desc, create_date desc', limit=1)
        return leads[:1]

    def _vd_qa_enrich_existing(self, core, info):
        """KH TRÙNG SĐT nhưng khách cũ CÒN THIẾU thông tin → điền nốt từ info file.
        CHỈ điền trường đang TRỐNG (không ghi đè dữ liệu đã có). True nếu có cập nhật."""
        if not info:
            return False
        vals = self._vd_qa_parse_info(info)
        if not vals:
            return False
        lead = self._vd_qa_find_lead_by_core(core)
        if not lead:
            return False
        write = {f: val for f, val in vals.items() if not lead[f]}
        if not write:
            return False
        lead.write(write)
        return True

    # ========================================================= PARSE INFO
    def _vd_qa_parse_info(self, text):
        """Parse 'Thông tin' → dict vd_intake_*. Luôn giữ nguyên văn ở Ghi chú.
        Nhận: diện tích (5x20 / 120m2), số tầng (2 tầng / T1,T2,T3 / cấp 4),
        diện tích từng tầng (T1: 90m2), mẫu nhà (mái bằng/thái/nhật), ngân sách."""
        t = (text or '').strip()
        if not t:
            return {}
        v = {'vd_intake_function_notes': t}
        nov = _novn(t)

        # Diện tích tổng: Rộng x Dài, hoặc "120m2".
        m = re.search(r'(\d+(?:[.,]\d+)?)\s*m?\s*x\s*(\d+(?:[.,]\d+)?)\s*m?', nov)
        if m:
            try:
                w = float(m.group(1).replace(',', '.'))
                l = float(m.group(2).replace(',', '.'))
                if 2 <= w <= 100 and 2 <= l <= 300:
                    v['vd_intake_area_m2'] = round(w * l, 1)
            except ValueError:
                pass
        if 'vd_intake_area_m2' not in v:
            m = re.search(r'(\d{2,4})\s*m2?\b', nov)
            if m and 10 <= float(m.group(1)) <= 5000:
                v['vd_intake_area_m2'] = float(m.group(1))

        # Số tầng: "2 tang" / "cap 4".
        fm = re.search(r'(\d)\s*tang', nov)
        if fm and fm.group(1) in '1234567':
            v['vd_intake_floors_select'] = fm.group(1)
        elif re.search(r'c[aâ]p\s*4', nov):
            v['vd_intake_floors_select'] = '1'

        # T1/T2/T3... = số tầng (viết tắt) + diện tích từng tầng "T1: 90m2".
        hits = re.findall(r'\bt\s*([1-7])\s*[:=]?\s*(\d{2,4})?\s*m?2?', nov)
        if hits:
            if 'vd_intake_floors_select' not in v:
                v['vd_intake_floors_select'] = str(max(int(f) for f, _a in hits))
            for fno, area in hits:
                if area and 10 <= float(area) <= 2000:
                    v['vd_intake_floor_%s_m2' % fno] = float(area)

        # Mẫu nhà (mái).
        ht = 'mai_thai' if 'mai thai' in nov else \
             'mai_nhat' if 'mai nhat' in nov else \
             'mai_bang' if 'mai bang' in nov else None
        if ht and self._vd_qa_sel_ok('vd_intake_house_type', ht):
            v['vd_intake_house_type'] = ht

        # Ngân sách: "2 ty" / "900 tr".
        am = re.search(r'(\d+(?:[.,]\d+)?)\s*(ty|ti|tr|trieu|k|nghin)', nov)
        if am:
            num = float(am.group(1).replace(',', '.'))
            unit = am.group(2)
            mult = 1e9 if unit in ('ty', 'ti') else 1e6 if unit in ('tr', 'trieu') else 1e3
            v['vd_intake_budget_amount'] = num * mult

        if v.get('vd_intake_floors_select') and any(
                v.get('vd_intake_floor_%d_m2' % i) for i in range(1, 8)):
            pass
        return v

    def _vd_qa_sel_ok(self, fname, key):
        fld = self.env['crm.lead']._fields.get(fname)
        if not fld or fld.type != 'selection':
            return False
        sel = fld.selection
        if callable(sel):
            sel = sel(self.env['crm.lead'])
        return key in {k for k, _l in (sel or [])}

    # ========================================================= EXCEL PARSE
    @api.model
    def vd_quickadd_parse_excel(self, file_b64, filename):
        """Đọc Excel/CSV → {rows (CHỈ số không trùng), dup, bad, total}.
        File cần cột SĐT; nhận thêm Tên / Nguồn / Thông tin nếu có."""
        try:
            data = base64.b64decode(file_b64)
        except Exception:
            raise UserError(_('File không hợp lệ.'))
        raw = self._vd_qa_read_rows(data, filename or '')
        _logger.info('[QUICKADD-DBG] file=%s | bytes=%d | raw_rows=%d | head=%s',
                     filename, len(data), len(raw),
                     [r[:6] for r in raw[:3]])
        if not raw:
            raise UserError(_('Không đọc được dòng nào. File cần có cột SĐT.'))
        rows = self._vd_qa_rows_to_dicts(raw)
        _logger.info('[QUICKADD-DBG] sdt_parse=%d | phones=%s',
                     len(rows), [r.get('phone') for r in rows])
        if not rows:
            raise UserError(_('Không tìm thấy SĐT nào trong file.'))
        for r in rows:
            r['excel'] = True
        clean, dups, bad = self._vd_qa_dedup(rows)
        _logger.info(
            '[QUICKADD-DBG] file=%s | raw_rows=%d | sdt_parse=%d | clean=%d dup=%d bad=%d'
            ' | phones=%s | dup_cores=%s',
            filename, len(raw), len(rows), len(clean), len(dups), bad,
            [r.get('phone') for r in rows], [d.get('core') for d in dups])
        # KH TRÙNG SĐT nhưng khách cũ CÒN THIẾU thông tin → bổ sung nốt từ file.
        enriched = 0
        for d in dups:
            if d.get('info') and self._vd_qa_enrich_existing(d['core'], d['info']):
                enriched += 1
        return {
            'rows': [{'phone': c['phone'], 'name': c['name'],
                      'source': c['source'], 'info': c['info'], 'excel': True}
                     for c in clean],
            'dup': len(dups), 'bad': bad, 'enriched': enriched, 'total': len(rows),
        }

    def _vd_qa_read_rows(self, data, filename):
        """bytes → list[list[str]]."""
        name = (filename or '').lower()
        if name.endswith('.csv') or (not name.endswith(('.xlsx', '.xls')) and b',' in data[:200]):
            text = None
            for enc in ('utf-8-sig', 'utf-8', 'cp1258', 'latin-1'):
                try:
                    text = data.decode(enc)
                    break
                except Exception:
                    continue
            if text is None:
                return []
            import csv
            delim = ';' if text.count(';') > text.count(',') else ','
            return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]
        # xlsx
        try:
            from openpyxl import load_workbook
        except Exception:
            raise UserError(_('Server thiếu thư viện openpyxl để đọc .xlsx.'))
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        out = []
        for row in ws.iter_rows(values_only=True):
            out.append(['' if c is None else self._vd_qa_cell(c) for c in row])
        return out

    @staticmethod
    def _vd_qa_cell(c):
        if isinstance(c, float) and c.is_integer():
            return str(int(c))
        return str(c).strip()

    def _vd_qa_rows_to_dicts(self, raw):
        """list[list] → list[{phone,name,source,info}].
        Tự tìm DÒNG HEADER (có thể KHÔNG phải dòng đầu — vd dòng đầu là ngày),
        khớp cột theo TỪ KHÓA CHỨA trong tên cột (nên 'Tên khách hàng',
        'SĐT', 'Thông tin', 'Nguồn' đều nhận ra); không thấy header → đoán."""
        PH = ('sdt', 'dien thoai', 'phone')
        NM = ('ten', 'ho ten', 'khach hang')
        SR = ('nguon', 'kenh', 'source')
        IF = ('thong tin', 'nhu cau', 'ghi chu', 'noi dung', 'info')

        def _match(h, kws):
            return any(k in h for k in kws)

        # Dòng header = dòng đầu tiên (trong 8 dòng đầu) có cột SĐT bằng CHỮ.
        hidx = -1
        header_norm = []
        for i, row in enumerate(raw[:8]):
            hn = [_novn(self._vd_qa_cell(c)).strip() for c in row]
            if any(_match(h, PH) for h in hn):
                hidx, header_norm = i, hn
                break

        ip = iname = isrc = iinfo = -1
        if hidx >= 0:
            def _first(kws):
                for j, h in enumerate(header_norm):
                    if _match(h, kws):
                        return j
                return -1
            ip, iname, isrc, iinfo = _first(PH), _first(NM), _first(SR), _first(IF)
            body = raw[hidx + 1:]
        else:
            body = raw

        out = []
        for row in body:
            cells = [self._vd_qa_cell(c) for c in row]
            if not any(cells):
                continue
            phone = name = source = info = ''
            if ip >= 0:
                phone = cells[ip] if ip < len(cells) else ''
                name = cells[iname] if 0 <= iname < len(cells) else ''
                source = cells[isrc] if 0 <= isrc < len(cells) else ''
                info = cells[iinfo] if 0 <= iinfo < len(cells) else ''
            else:
                # Không header: tìm ô giống SĐT, ô chữ dài nhất làm tên.
                for c in cells:
                    digits = re.sub(r'\D', '', c)
                    if 9 <= len(digits) <= 12 and not phone:
                        phone = c
                    elif re.search(r'[A-Za-zÀ-ỹ]', c) and len(c) > len(name):
                        name = c
            if not re.sub(r'\D', '', phone):
                continue
            out.append({'phone': phone, 'name': name,
                        'source': self._vd_qa_norm_source(source), 'info': info})
        return out

    @staticmethod
    def _vd_qa_norm_source(txt):
        s = _novn(txt)
        if 'zalo' in s:
            return 'zalo'
        if 'tiktok' in s or 'tik tok' in s:
            return 'tiktok'
        if 'face' in s or s == 'fb':
            return 'facebook'
        return 'quet'

    # =========================================================== DISTRIBUTE
    def _vd_qa_target_order(self, nv_ids):
        """Recordset NV đích cho chia số. CHIA ĐỀU TRONG NGÀY: ưu tiên NV được
        chia ÍT KHÁCH NHẤT HÔM NAY (mọi nguồn) → nhiều lần upload nhỏ KHÔNG dồn
        hết vào NV ít tổng khách nhất. Hoà → tải khách mới ít hơn, rồi id nhỏ hơn.
        nv_ids rỗng → mọi NV đang nhận số."""
        from datetime import datetime, time as _t, timedelta as _td
        import pytz
        Users = self.env['res.users'].sudo()
        if nv_ids:
            pool = Users.browse(nv_ids).exists()
        else:
            pool = Users.browse([d['id'] for d in self._vd_qa_nv_list()]).filtered(
                lambda u: bool(getattr(u, 'vd_can_receive_pancake', True)))
        if not pool:
            raise UserError(_('Chưa chọn nhân viên nào để chia số.'))
        new_stage = self.env.ref('vd_crm_lead.stage_new', raise_if_not_found=False)
        Lead = self.env['crm.lead'].sudo()
        # Số KH đã chia HÔM NAY (giờ VN) cho mỗi NV — mọi nguồn (quét + Pancake).
        vn = pytz.timezone('Asia/Ho_Chi_Minh')
        d0 = datetime.now(vn).date()
        day_s = vn.localize(datetime.combine(d0, _t(0, 0))).astimezone(
            pytz.utc).replace(tzinfo=None)
        day_e = day_s + _td(days=1)
        today = {}
        for g in Lead.read_group(
                [('user_id', 'in', pool.ids),
                 ('create_date', '>=', day_s), ('create_date', '<', day_e)],
                ['user_id'], ['user_id']):
            if g.get('user_id'):
                today[g['user_id'][0]] = g.get('user_id_count', 0)
        load = {}
        for u in pool:
            dom = [('user_id', '=', u.id), ('active', '=', True)]
            if new_stage:
                dom.append(('stage_id', '=', new_stage.id))
            load[u.id] = Lead.search_count(dom)
        return pool.sorted(
            lambda u: (today.get(u.id, 0), load.get(u.id, 0), u.id))

    @api.model
    def vd_quickadd_distribute(self, lead_ids, nv_ids=None):
        """Chia số KH ĐÃ CÓ (chọn bằng kéo chuột) cho NV — chỉ đổi NV phụ trách,
        GIỮ nguyên dữ liệu. nv_ids rỗng → chia đều NV đang nhận số.
        Round-robin theo tải. Trả {moved, detail}."""
        leads = self.env['crm.lead'].sudo().browse(lead_ids or []).exists()
        if not leads:
            raise UserError(_('Chưa chọn khách nào để chia.'))
        order = self._vd_qa_target_order(nv_ids)
        n = len(order)
        Users = self.env['res.users'].sudo()
        from collections import Counter
        counts = Counter()
        for i, lead in enumerate(leads):
            u = order[i % n]
            lead.with_context(
                vd_skip_reassign_check=True, vd_skip_assignment_balance=True,
            ).write({'user_id': u.id})
            counts[u.id] += 1
        detail = ', '.join('%s (%d)' % (Users.browse(uid).name or '?', c)
                           for uid, c in counts.most_common())
        return {'moved': len(leads), 'detail': detail}

    # ============================================================== SUBMIT
    @api.model
    def vd_quickadd_submit(self, rows, nv_ids=None):
        """Tạo lead từ rows + chia cho NV. rows=[{phone,name,source,info,excel}].
        nv_ids rỗng → chia đều cho NV đang nhận số. Trả {created, dup, bad, detail}."""
        rows = rows or []
        if not rows:
            raise UserError(_('Chưa có khách nào để tạo.'))
        clean, dups, bad = self._vd_qa_dedup(rows)
        dup = len(dups)
        # KH trùng nhưng có info → bổ sung thông tin còn thiếu cho khách cũ.
        enriched = 0
        for d in dups:
            if d.get('info') and self._vd_qa_enrich_existing(d['core'], d['info']):
                enriched += 1
        if not clean:
            if enriched:
                return {'created': 0, 'dup': dup, 'bad': bad, 'enriched': enriched,
                        'detail': '', 'only_enriched': True}
            raise UserError(_('Tất cả %d số đều TRÙNG hoặc SAI — không có số mới.') % len(rows))

        order = self._vd_qa_target_order(nv_ids)
        n = len(order)
        new_stage = self.env.ref('vd_crm_lead.stage_new', raise_if_not_found=False)
        Lead = self.env['crm.lead'].sudo()
        Users = self.env['res.users'].sudo()

        from collections import Counter
        counts = Counter()
        created = Lead.browse()
        for i, c in enumerate(clean):
            u = order[i % n]
            channel = VD_QA_CHANNEL.get(c['source'], 'other')
            vals = {
                'name': c['name'], 'partner_name': c['name'], 'phone': c['phone'],
                'user_id': u.id, 'type': 'lead',
                'vd_lead_channel': channel,
                # Mọi số thêm qua popup (gõ tay HOẶC Excel) đều là "QUÉT SỐ" —
                # vd_from_excel là cờ "quét số" cho báo cáo chia số (không để số
                # quét lọt vào cột "tự động" của Pancake) (fix 2026-10-08).
                'vd_from_excel': True,
            }
            if new_stage:
                vals['stage_id'] = new_stage.id
            if c['info']:
                vals.update(self._vd_qa_parse_info(c['info']))
            lead = Lead.with_context(
                vd_skip_reassign_check=True, vd_skip_assignment_balance=True,
            ).create(vals)
            created |= lead
            counts[u.id] += 1
        detail = ', '.join('%s (%d)' % (Users.browse(uid).name or '?', c)
                           for uid, c in counts.most_common())
        return {'created': len(created), 'dup': dup, 'bad': bad,
                'enriched': enriched, 'detail': detail}
